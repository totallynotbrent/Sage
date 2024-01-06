"""PPTX extraction via python-pptx: one unit per slide (concatenated shapes)."""

from __future__ import annotations

from io import BytesIO

from app.services.extraction.base import (
    ExtractionResult,
    ExtractedUnit,
    LocationInfo,
    register,
    register_unavailable,
)

try:
    from pptx import Presentation

    EXTRACTION_AVAILABLE = True
except ImportError:
    EXTRACTION_AVAILABLE = False

if not EXTRACTION_AVAILABLE:
    register_unavailable("pptx", "python-pptx")
else:

    @register("pptx")
    class PptxExtractor:
        """Concatenate the text of every shape on each slide."""

        def extract(self, data: bytes, *, filename: str) -> ExtractionResult:
            try:
                presentation = Presentation(BytesIO(data))
            except Exception as exc:  # noqa: BLE001 - convert to result
                return ExtractionResult(error=f"Could not open PPTX: {exc}")

            units: list[ExtractedUnit] = []
            warnings: list[str] = []
            for slide_index, slide in enumerate(presentation.slides, start=1):
                parts: list[str] = []
                for shape in slide.shapes:
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    text = (shape.text_frame.text or "").strip()
                    if text:
                        parts.append(text)
                text = "\n".join(parts).strip()
                if not text:
                    warnings.append(f"slide {slide_index} contains no text")
                    continue
                units.append(
                    ExtractedUnit(
                        text=text,
                        location=LocationInfo(kind="slide", slide=slide_index),
                    )
                )

            if not units:
                return ExtractionResult(
                    error="No extractable text found in this PPTX.",
                    warnings=warnings,
                )
            return ExtractionResult(units=units, warnings=warnings)
